from datetime import date
from io import BytesIO

from django.test import TestCase
from openpyxl import load_workbook

from api.models import EmployeePfEsiDetail, EmployeeSalaryPrepared
from api.reports.pf_esi_reports.generate_esi_statement_xlsx import generate_esi_statement_xlsx
from api.tests.base import AttendanceTestDataMixin


class EsiStatementReportTests(AttendanceTestDataMixin, TestCase):
    def test_statement_uses_prepared_salary_esi_snapshots(self):
        employee = self.create_employee()
        EmployeePfEsiDetail.objects.create(
            user=self.user,
            company=self.company,
            employee=employee,
            esi_allow=True,
            esi_number='ESI-1',
        )
        EmployeeSalaryPrepared.objects.create(
            user=self.user,
            company=self.company,
            employee=employee,
            date=date(2024, 1, 1),
            esi_employee_wages=1234,
            esi_deducted=10,
            esi_employer_contribution=40,
        )

        response = generate_esi_statement_xlsx(
            self.user,
            {'year': 2024, 'month': 1},
            [employee],
        )
        worksheet = load_workbook(BytesIO(response.content), data_only=True).active

        self.assertEqual(worksheet.cell(row=2, column=6).value, 1234)
        self.assertEqual(worksheet.cell(row=2, column=7).value, 10)
        self.assertEqual(worksheet.cell(row=2, column=8).value, 40)
