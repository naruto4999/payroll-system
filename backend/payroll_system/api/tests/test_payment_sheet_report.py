from datetime import date
from io import BytesIO

from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook
from rest_framework.test import APIClient

from api.models import (
    CompanyReportConfiguration,
    Deparment,
    EarnedAmount,
    EmployeeMonthlyAttendanceDetails,
    PfEsiSetup,
)
from api.reports.payment_sheet_data import build_payment_sheet_report
from api.tests.base import AttendanceTestDataMixin


def payment_sheet_headers(worksheet):
    return [
        worksheet.cell(row=2, column=column).value
        or worksheet.cell(row=1, column=column).value
        for column in range(1, worksheet.max_column + 1)
    ]


class PaymentSheetReportTests(AttendanceTestDataMixin, TestCase):
    def setUp(self):
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        company_details = self.company.company_details
        company_details.address = 'Test address'
        company_details.save(update_fields=['address'])

    def create_report_employee(self, index):
        employee = self.create_employee(
            paycode=f'E{index:03d}',
            attendance_card_no=100 + index,
        )
        salary_rate = self.create_salary_earning(employee, value=20800 + index)
        EmployeeMonthlyAttendanceDetails.objects.create(
            user=self.user,
            company=self.company,
            employee=employee,
            date=date(2024, 1, 1),
            present_count=52,
            paid_days_count=61,
        )
        prepared_salary = self.create_prepared_salary(
            employee,
            net_minutes=90,
            amount=300,
        )
        prepared_salary.incentive_amount = 100
        prepared_salary.advance_deducted = 10
        prepared_salary.pf_deducted = 20
        prepared_salary.vpf_deducted = 5
        prepared_salary.esi_deducted = 6
        prepared_salary.tds_deducted = 7
        prepared_salary.others_deducted = 8
        prepared_salary.save(update_fields=[
            'incentive_amount',
            'advance_deducted',
            'pf_deducted',
            'vpf_deducted',
            'esi_deducted',
            'tds_deducted',
            'others_deducted',
        ])
        EarnedAmount.objects.create(
            user=self.user,
            earnings_head=salary_rate.earnings_head,
            salary_prepared=prepared_salary,
            rate=salary_rate.value,
            earned_amount=20000 + index,
        )
        return employee

    def payload(self, employees, output_format, group_by='none', **filter_overrides):
        return {
            'report_type': 'payment_sheet',
            'employee_ids': [employee.pk for employee in employees],
            'company': self.company.pk,
            'month': 1,
            'year': 2024,
            'filters': {
                'group_by': group_by,
                'resignation_filter': 'all',
                'sort_by': 'attendance_card_no',
                'language': 'english',
                'format': output_format,
                'overtime': 'with_ot',
                **filter_overrides,
            },
        }

    def render(self, employees, output_format, group_by='none', **filter_overrides):
        with CaptureQueriesContext(connection) as queries:
            response = self.client.post(
                '/api/generate-salary-overtime-sheet',
                self.payload(employees, output_format, group_by, **filter_overrides),
                format='json',
            )
            content = b''.join(response.streaming_content)
        return response, content, queries.captured_queries

    def test_xlsx_contains_employee_values_and_totals(self):
        employee = self.create_report_employee(1)

        response, content, _ = self.render([employee], 'xlsx')

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        workbook = load_workbook(BytesIO(content), data_only=True)
        worksheet = workbook['Sheet1']
        headers = payment_sheet_headers(worksheet)
        employee_values = dict(zip(headers, [cell.value for cell in worksheet[3]]))
        total_values = dict(zip(headers, [cell.value for cell in worksheet[4]]))

        self.assertEqual(employee_values['ACN'], 101)
        self.assertEqual(employee_values['Salary Rate'], 20801)
        self.assertEqual(employee_values['PD'], 30.5)
        self.assertEqual(employee_values['Earned Salary (*incl of arrear)'], 20001)
        self.assertEqual(employee_values['OT Hrs'], 1.5)
        self.assertEqual(employee_values['Total Earned'], 20401)
        self.assertEqual(employee_values['EPF'], 25)
        self.assertEqual(employee_values['Total Deduction'], 56)
        self.assertEqual(employee_values['Net Payable'], 20345)
        self.assertEqual(total_values['S/N'], 'Grand Total')
        self.assertEqual(total_values['Net Payable'], 20345)

    def test_grouped_xlsx_contains_department_totals_and_lwf(self):
        department = Deparment.objects.create(
            user=self.user,
            company=self.company,
            name='Operations',
        )
        employee = self.create_report_employee(1)
        professional_detail = employee.employee_professional_detail
        professional_detail.department = department
        professional_detail.save(update_fields=['department'])
        prepared_salary = employee.salaries_prepared.get(date=date(2024, 1, 1))
        prepared_salary.labour_welfare_fund_deducted = 9
        prepared_salary.save(update_fields=['labour_welfare_fund_deducted'])
        setup = PfEsiSetup.objects.get(company=self.company)
        setup.enable_labour_welfare_fund = True
        setup.save(update_fields=['enable_labour_welfare_fund'])

        response, content, _ = self.render([employee], 'xlsx', group_by='department')

        self.assertEqual(response.status_code, 200)
        worksheet = load_workbook(BytesIO(content), data_only=True)['Sheet1']
        headers = payment_sheet_headers(worksheet)
        self.assertIn('LWF', headers)
        self.assertEqual(worksheet.cell(row=3, column=1).value, 'Operations')
        self.assertEqual(worksheet.cell(row=5, column=1).value, 'Department Total')
        department_total = dict(zip(headers, [cell.value for cell in worksheet[5]]))
        self.assertEqual(department_total['LWF'], 9)
        self.assertEqual(department_total['Total Deduction'], 65)

        pdf_response, pdf_content, _ = self.render(
            [employee], 'pdf', group_by='department'
        )
        self.assertEqual(pdf_response.status_code, 200)
        self.assertTrue(pdf_content.startswith(b'%PDF-'))

    def test_xlsx_can_show_salary_and_earned_amounts_by_head(self):
        employee = self.create_report_employee(1)
        hra_rate = self.create_salary_earning(employee, name='HRA', value=5000)
        prepared_salary = employee.salaries_prepared.get(date=date(2024, 1, 1))
        EarnedAmount.objects.create(
            user=self.user,
            earnings_head=hra_rate.earnings_head,
            salary_prepared=prepared_salary,
            rate=hra_rate.value,
            earned_amount=4500,
        )

        response, content, _ = self.render(
            [employee],
            'xlsx',
            salary_rate_columns='head_wise',
            earned_salary_columns='head_wise',
        )

        self.assertEqual(response.status_code, 200)
        worksheet = load_workbook(BytesIO(content), data_only=True)['Sheet1']
        headers = payment_sheet_headers(worksheet)
        employee_values = dict(zip(headers, [cell.value for cell in worksheet[3]]))
        total_values = dict(zip(headers, [cell.value for cell in worksheet[4]]))
        self.assertEqual(headers[5:12], [
            'Salary Rate - Basic',
            'Salary Rate - HRA',
            'Salary Rate Total',
            'PD',
            'Earned - Basic (*incl of arrear)',
            'Earned - HRA (*incl of arrear)',
            'Earned Salary Total (*incl of arrear)',
        ])
        self.assertEqual(worksheet['F1'].value, 'SALARY RATE')
        self.assertEqual(worksheet['J1'].value, 'EARNED SALARY')
        self.assertIn('F1:H1', {str(cell_range) for cell_range in worksheet.merged_cells})
        self.assertIn('J1:L1', {str(cell_range) for cell_range in worksheet.merged_cells})
        self.assertEqual(worksheet.freeze_panes, 'F3')
        self.assertTrue(worksheet['F1'].fill.fgColor.rgb.endswith('C97D7D'))
        self.assertTrue(worksheet['J1'].fill.fgColor.rgb.endswith('3D85C6'))
        self.assertEqual(employee_values['Salary Rate Total'], 25801)
        self.assertEqual(employee_values['Salary Rate - Basic'], 20801)
        self.assertEqual(employee_values['Salary Rate - HRA'], 5000)
        self.assertEqual(employee_values['Earned Salary Total (*incl of arrear)'], 24501)
        self.assertEqual(employee_values['Earned - Basic (*incl of arrear)'], 20001)
        self.assertEqual(employee_values['Earned - HRA (*incl of arrear)'], 4500)
        self.assertEqual(total_values['Salary Rate - HRA'], 5000)
        self.assertEqual(total_values['Earned - HRA (*incl of arrear)'], 4500)

    def test_company_configuration_can_show_weighted_ot_hours(self):
        employee = self.create_report_employee(1)
        prepared_salary = employee.salaries_prepared.get(date=date(2024, 1, 1))
        self.create_prepared_overtime_detail(
            prepared_salary,
            day_type='REGULAR',
            gross_minutes=60,
            net_minutes=60,
            multiplier='1.5',
        )
        self.create_prepared_overtime_detail(
            prepared_salary,
            day_type='HOLIDAY',
            gross_minutes=30,
            net_minutes=30,
            multiplier='2',
        )
        CompanyReportConfiguration.objects.create(
            company=self.company,
            report_type='payment_sheet',
            output_format='xlsx',
            options={'ot_hours_display': 'weighted'},
        )

        response, content, _ = self.render([employee], 'xlsx')

        self.assertEqual(response.status_code, 200)
        worksheet = load_workbook(BytesIO(content), data_only=True)['Sheet1']
        headers = payment_sheet_headers(worksheet)
        employee_values = dict(zip(headers, [cell.value for cell in worksheet[3]]))
        total_values = dict(zip(headers, [cell.value for cell in worksheet[4]]))
        self.assertNotIn('Weighted OT Hrs', headers)
        self.assertEqual(employee_values['OT Hrs'], 2.5)
        self.assertEqual(total_values['OT Hrs'], 2.5)
        self.assertEqual(employee_values['OT Amount'], 300)

    def test_weighted_ot_hours_fall_back_to_actual_hours_without_snapshot(self):
        employee = self.create_report_employee(1)
        CompanyReportConfiguration.objects.create(
            company=self.company,
            report_type='payment_sheet',
            output_format='xlsx',
            options={'ot_hours_display': 'weighted'},
        )

        _, content, _ = self.render([employee], 'xlsx')

        worksheet = load_workbook(BytesIO(content), data_only=True)['Sheet1']
        headers = payment_sheet_headers(worksheet)
        employee_values = dict(zip(headers, [cell.value for cell in worksheet[3]]))
        self.assertEqual(employee_values['OT Hrs'], 1.5)

    def test_pdf_and_xlsx_query_counts_do_not_scale_per_employee(self):
        employees = [self.create_report_employee(1)]
        pdf_response, pdf_content, one_pdf_queries = self.render(employees, 'pdf')
        xlsx_response, xlsx_content, one_xlsx_queries = self.render(employees, 'xlsx')

        self.assertEqual(pdf_response.status_code, 200)
        self.assertTrue(pdf_content.startswith(b'%PDF-'))
        self.assertEqual(xlsx_response.status_code, 200)
        self.assertTrue(xlsx_content.startswith(b'PK'))

        employees.extend(self.create_report_employee(index) for index in range(2, 6))
        _, _, five_pdf_queries = self.render(employees, 'pdf')
        _, _, five_xlsx_queries = self.render(employees, 'xlsx')

        self.assertLessEqual(
            len(five_pdf_queries),
            len(one_pdf_queries) + 1,
            [query['sql'] for query in five_pdf_queries],
        )
        self.assertLessEqual(
            len(five_xlsx_queries),
            len(one_xlsx_queries) + 1,
            [query['sql'] for query in five_xlsx_queries],
        )
        self.assertLessEqual(len(five_pdf_queries), 12)
        self.assertLessEqual(len(five_xlsx_queries), 12)

    def test_format_specific_missing_attendance_behavior_is_preserved(self):
        employee = self.create_report_employee(1)
        EmployeeMonthlyAttendanceDetails.objects.filter(employee=employee).delete()
        salaries = list(
            employee.salaries_prepared.select_related(
                'employee',
                'employee__employee_professional_detail',
                'employee__employee_professional_detail__department',
                'employee__employee_professional_detail__designation',
                'company',
                'company__company_details',
                'user',
                'user__regular_to_owner',
            )
        )
        request_data = self.payload([employee], 'pdf')

        pdf_row = build_payment_sheet_report(
            self.user, request_data, salaries, 'pdf'
        ).rows[0]
        xlsx_row = build_payment_sheet_report(
            self.user, request_data, salaries, 'xlsx'
        ).rows[0]

        self.assertIsNone(pdf_row.paid_days)
        self.assertIsNone(pdf_row.ot_hours)
        self.assertIsNone(pdf_row.ot_amount)
        self.assertEqual(pdf_row.total_earned, 20101)
        self.assertIsNone(xlsx_row.paid_days)
        self.assertEqual(xlsx_row.ot_hours, 1.5)
        self.assertEqual(xlsx_row.ot_amount, 300)
        self.assertEqual(xlsx_row.total_earned, 20401)
